"""Checks that stop execution, and malformed inputs that must not pass.

Each test here pins down a behavior the S001 metapaper states: what raises,
what reports, and which inputs are refused rather than silently repaired.
"""

import subprocess

import pytest


# --- provenance -----------------------------------------------------------

def _manifest(tmp_path, files):
    from edikit.data.provenance import write_manifest
    paths = []
    for name, body in files.items():
        p = tmp_path / name
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(body)
        paths.append(p)
    m = tmp_path / "SOURCES.sha256"
    write_manifest(paths, m)
    return m


def test_verify_sources_passes_on_registered_files(tmp_path):
    from edikit.data.provenance import verify_sources
    m = _manifest(tmp_path, {"raw/a.csv": b"x,1\n", "raw/b.csv": b"y,2\n"})
    res = verify_sources(m)
    assert res.ok and sorted(res.verified) == ["raw/a.csv", "raw/b.csv"]


def test_verify_sources_raises_on_altered_file(tmp_path):
    from edikit.data.provenance import verify_sources
    m = _manifest(tmp_path, {"raw/a.csv": b"x,1\n"})
    (tmp_path / "raw/a.csv").write_bytes(b"x,2\n")       # a different vintage
    with pytest.raises(ValueError, match="differs from the manifest"):
        verify_sources(m)


def test_verify_sources_raises_on_missing_file(tmp_path):
    from edikit.data.provenance import verify_sources
    m = _manifest(tmp_path, {"raw/a.csv": b"x,1\n"})
    (tmp_path / "raw/a.csv").unlink()
    with pytest.raises(FileNotFoundError, match="registered source missing"):
        verify_sources(m)


def test_verify_sources_collect_mode_lists_every_problem(tmp_path):
    from edikit.data.provenance import verify_sources
    m = _manifest(tmp_path, {"a": b"1", "b": b"2", "c": b"3"})
    (tmp_path / "a").write_bytes(b"changed")
    (tmp_path / "b").unlink()
    res = verify_sources(m, raise_on_error=False)
    assert not res.ok
    assert res.verified == ["c"] and res.missing == ["b"]
    assert [r[0] for r in res.mismatched] == ["a"]


def test_manifest_rejects_malformed_and_duplicate_lines(tmp_path):
    from edikit.data.provenance import read_manifest
    bad = tmp_path / "bad.sha256"
    bad.write_text("not-a-hash  raw/a.csv\n")
    with pytest.raises(ValueError, match="not a sha256sum entry"):
        read_manifest(bad)
    dup = tmp_path / "dup.sha256"
    h = "0" * 64
    dup.write_text(f"{h}  raw/a.csv\n{h}  raw/a.csv\n")
    with pytest.raises(ValueError, match="more than once"):
        read_manifest(dup)


def test_verify_sources_only_must_name_registered_paths(tmp_path):
    from edikit.data.provenance import verify_sources
    m = _manifest(tmp_path, {"a": b"1"})
    with pytest.raises(ValueError, match="not registered"):
        verify_sources(m, only=["z"])


def test_manifest_is_readable_by_standard_tools(tmp_path):
    """The manifest format is sha256sum's, so `shasum -a 256 -c` verifies it
    with no Python involved. Skipped where neither tool is installed."""
    m = _manifest(tmp_path, {"raw/a.csv": b"x,1\n"})
    for cmd in (["sha256sum", "-c", m.name],
                ["shasum", "-a", "256", "-c", m.name]):
        try:
            out = subprocess.run(cmd, cwd=tmp_path, capture_output=True,
                                 text=True)
        except FileNotFoundError:
            continue
        assert out.returncode == 0, out.stdout + out.stderr
        return
    pytest.skip("neither sha256sum nor shasum is available")


# --- classification in the audit ----------------------------------------

_CELLS = {("a", "b"): 5.0, ("b", "a"): 5.0}


def test_audit_default_reports_unclassified_and_phantom_accounts():
    from edikit.pipeline.audit import audit
    res = audit(_CELLS, kinds={"a": "activity", "ghost": "household"})
    assert res.coverage_gap == ["b"]      # in the matrix, no kind
    assert res.phantom == ["ghost"]       # has a kind, not in the matrix


def test_audit_strict_raises_on_unclassified_account():
    from edikit.pipeline.audit import audit
    with pytest.raises(ValueError, match="classification incomplete"):
        audit(_CELLS, kinds={"a": "activity"}, strict=True)


def test_audit_strict_raises_on_phantom_account():
    from edikit.pipeline.audit import audit
    with pytest.raises(ValueError, match="classification incomplete"):
        audit(_CELLS, kinds={"a": "activity", "b": "commodity",
                             "ghost": "household"}, strict=True)


def test_audit_strict_passes_on_complete_classification():
    from edikit.pipeline.audit import audit
    res = audit(_CELLS, kinds={"a": "activity", "b": "commodity"},
                strict=True)
    assert not res.coverage_gap and not res.phantom


def test_read_kinds_rejects_unknown_kind_as_error_not_assert(tmp_path):
    """Formerly a bare assert, which `python -O` strips."""
    from edikit.pipeline.audit import read_kinds
    p = tmp_path / "kinds.csv"
    p.write_text("account,kind\na,activity\nb,not-a-kind\n")
    with pytest.raises(ValueError, match="unknown kind"):
        read_kinds(p)


# --- duplicate identifiers: refused, never overwritten silently ----------

def test_long_csv_rejects_repeated_cell(tmp_path):
    from edikit.pipeline.audit import read_long_csv
    p = tmp_path / "sam.csv"
    p.write_text("row,col,value\na,b,1\na,b,2\n")
    with pytest.raises(ValueError, match="appears more than once"):
        read_long_csv(p)


def test_matrix_csv_rejects_repeated_row_code(tmp_path):
    from edikit.pipeline.audit import read_matrix_csv
    p = tmp_path / "sam.csv"
    p.write_text(",a,b\na,0,1\nb,1,0\na,0,2\n")
    with pytest.raises(ValueError, match="row code 'a'"):
        read_matrix_csv(p)


def test_matrix_csv_rejects_repeated_column_code(tmp_path):
    from edikit.pipeline.audit import read_matrix_csv
    p = tmp_path / "sam.csv"
    p.write_text(",a,a\na,0,1\nb,1,0\n")
    with pytest.raises(ValueError, match="column codes"):
        read_matrix_csv(p)


def test_matrix_csv_still_skips_total_row_and_column(tmp_path):
    from edikit.pipeline.audit import read_matrix_csv
    p = tmp_path / "sam.csv"
    p.write_text(",a,b,total\na,0,1,1\nb,1,0,1\ntotal,1,1,2\n")
    assert read_matrix_csv(p) == {("a", "b"): 1.0, ("b", "a"): 1.0}


def test_concordance_rejects_source_mapped_twice(tmp_path):
    from edikit.data.concordance import Concordance
    p = tmp_path / "c.csv"
    p.write_text("source,target\nX1,A\nX1,B\n")
    with pytest.raises(ValueError, match="mapped more than once"):
        Concordance.from_csv(str(p))


# --- named tolerances replace magic numbers ------------------------------

def test_tiling_tolerance_is_a_named_documented_constant():
    from edikit.pipeline import eurostat_sam
    assert eurostat_sam.TILING_TOLERANCE == 0.05


# --- output precision is stated and selectable ----------------------------

def test_balanced_csv_precision_is_selectable(tmp_path):
    """Default 0.1m keeps the committed files byte-identical; a larger
    `decimals` gives a higher-precision export of the same matrix."""
    from edikit.model.balancing import ras
    from edikit.pipeline.eurostat_sam import write_balanced_csv
    seed = {("a", "b"): 1.0, ("b", "a"): 2.0, ("a", "a"): 1.0}
    bal = ras(seed, {"a": 2.5, "b": 1.5}, {"a": 2.5, "b": 1.5})
    write_balanced_csv(bal, tmp_path / "lo.csv")
    write_balanced_csv(bal, tmp_path / "hi.csv", decimals=6)
    lo = (tmp_path / "lo.csv").read_text().splitlines()[1].split(",")
    hi = (tmp_path / "hi.csv").read_text().splitlines()[1].split(",")
    assert len(lo[2].split(".")[1]) == 1 and len(hi[2].split(".")[1]) == 6
    assert abs(float(lo[2]) - float(hi[2])) <= 0.05


# --- rounding is tested, not assumed --------------------------------------

def test_audit_rounding_bound_separates_rounding_from_real_imbalance(tmp_path):
    """An imbalance within half a grain per nonzero cell can be rounding;
    one beyond it cannot, and the report must say which accounts exceed."""
    from edikit.pipeline.audit import audit, write_report
    # every account has 4 nonzero cells (2 in its row, 2 in its column), so
    # each bound is 0.5 * 1 * 4 = 2. Gaps: a +1 (rounding can explain it),
    # b -11 and c +10 (it cannot).
    cells = {("a", "b"): 10.0, ("a", "c"): 10.0, ("b", "a"): 9.0,
             ("c", "a"): 10.0, ("c", "b"): 30.0, ("b", "c"): 20.0}
    res = audit(cells)
    assert res.grain == 1.0
    assert abs(res.gaps["a"]) <= res.rounding_bound["a"]
    assert abs(res.gaps["c"]) > res.rounding_bound["c"]
    write_report(res, tmp_path / "a.md")
    text = (tmp_path / "a.md").read_text()
    line = text.split("Rounding test")[1].splitlines()[0]
    assert line.rstrip().endswith("2/3 (b, c)")      # a is within its bound


def test_audit_balance_count_without_grain_uses_printed_precision(tmp_path):
    """Unrounded values leave floating-point gaps; an account balanced to the
    four decimals the report prints must count as balanced."""
    from edikit.pipeline.audit import audit, write_report
    cells = {("a", "b"): 1.23456, ("b", "a"): 1.23457,
             ("b", "c"): 2.5, ("c", "b"): 2.50001}
    res = audit(cells)
    assert res.grain == 0.0
    write_report(res, tmp_path / "a.md")
    text = (tmp_path / "a.md").read_text()
    assert "accounts balanced to the four decimals shown: 3/3" in text


# --- every refusal in Table 1 is exercised --------------------------------

def test_generate_raises_when_country_year_has_no_data(tmp_path):
    from conftest import synthetic_economy_paths
    from edikit.pipeline.eurostat_sam import generate
    paths = synthetic_economy_paths(tmp_path)
    with pytest.raises(ValueError, match="no supply-table data"):
        generate("YY", 2019, paths)          # fixture holds only XX


def test_generate_raises_on_double_counting(tmp_path):
    """Leaf industries exceeding the published total by >5% is a partition
    that double-counts; it must stop, not be balanced away."""
    from conftest import synthetic_economy_paths
    from edikit.pipeline.eurostat_sam import generate
    paths = synthetic_economy_paths(tmp_path,
                                    supply_edits={("A01", "CPA_A01"): 130})
    with pytest.raises(ValueError, match="double-counts"):
        generate("XX", 2019, paths)


def test_generate_raises_on_industry_absent_from_use_table(tmp_path):
    from conftest import synthetic_economy_paths
    from edikit.pipeline.eurostat_sam import generate
    paths = synthetic_economy_paths(tmp_path, drop_use_industry="B")
    with pytest.raises(ValueError, match="absent from use"):
        generate("XX", 2019, paths)


def test_ras_raises_on_positive_target_for_empty_column():
    from edikit.model.balancing import ras
    with pytest.raises(ValueError, match="empty seed column"):
        # both rows have seed entries; column b has none
        ras({("a", "a"): 1.0, ("b", "a"): 1.0},
            {"a": 1.0, "b": 1.0}, {"a": 1.0, "b": 1.0})


def test_ras_records_non_convergence_without_raising():
    from edikit.model.balancing import ras
    seed = {("a", "a"): 1.0, ("a", "b"): 1.0, ("b", "a"): 1.0}
    res = ras(seed, {"a": 5.0, "b": 1.0}, {"a": 1.0, "b": 5.0}, max_iter=2)
    assert res.converged is False


def test_concordance_aggregate_raises_on_unmapped_code():
    from edikit.data.concordance import Concordance
    c = Concordance(name="t", mapping={"X1": "A"})
    with pytest.raises(KeyError):
        c.aggregate({"X1": 1.0, "X2": 2.0})


# --- residual transparency (cross-source check, items, notes) -------------

def test_generate_records_vintages_cross_source_and_residual_items(tmp_path):
    from conftest import synthetic_economy_paths
    from edikit.pipeline.eurostat_sam import balance, generate, write_report
    res = generate("XX", 2019, synthetic_economy_paths(tmp_path))
    assert set(res.vintages) == {"supply", "use", "sectors"}
    assert "Gross value added at basic prices" in res.cross_source
    assert ("row", "cap") in res.plugs and "corp" in res.residual_items
    bal, flipped = balance(res)
    write_report(res, tmp_path / "r.md", balance=bal, flipped=flipped)
    text = (tmp_path / "r.md").read_text()
    for section in ("## Cross-source consistency",
                    "## Items bearing on each residual",
                    "## Closing entries", "capital factor account"):
        assert section in text
    assert "ESA boundary items" not in text      # the old, unsupported claim


def test_balanced_csv_marks_closing_and_flipped_cells(tmp_path):
    from edikit.model.balancing import ras
    from edikit.pipeline.eurostat_sam import write_balanced_csv
    seed = {("a", "b"): 1.0, ("b", "a"): 2.0, ("a", "a"): 1.0}
    bal = ras(seed, {"a": 2.5, "b": 1.5}, {"a": 2.5, "b": 1.5})
    write_balanced_csv(bal, tmp_path / "plain.csv")
    assert "note" not in (tmp_path / "plain.csv").read_text().splitlines()[0]
    write_balanced_csv(bal, tmp_path / "noted.csv", flipped=[("a", "b")],
                       plugs={("a", "b"): -1.0})
    rows = (tmp_path / "noted.csv").read_text().splitlines()
    assert rows[0].endswith(",note")
    assert any("closing entry, flipped from (a,b)" in r for r in rows)


# --- the readers that still kept the last duplicate ----------------------

def test_read_kinds_rejects_account_classified_twice(tmp_path):
    from edikit.pipeline.audit import read_kinds
    p = tmp_path / "kinds.csv"
    p.write_text("account,kind\na1,activity\na1,commodity\n")
    with pytest.raises(ValueError, match="classified more than once"):
        read_kinds(p)


def test_read_controls_rejects_repeated_control(tmp_path):
    from edikit.pipeline.audit import read_controls
    p = tmp_path / "controls.csv"
    p.write_text("aggregate,value\nExports,1\nExports,2\n")
    with pytest.raises(ValueError, match="appears more than once"):
        read_controls(p)


def test_labeled_matrix_rejects_repeated_labels(tmp_path):
    import openpyxl
    from edikit.data.sam import read_labeled_matrix
    for name, rows in [
            ("dup_row", [[None, "a", "b", "c"], ["a", 1, 2, 3],
                         ["a", 4, 5, 6]]),
            ("dup_col", [[None, "a", "a", "c"], ["a", 1, 2, 3]])]:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SAM"
        for r in rows:
            ws.append(r)
        p = tmp_path / f"{name}.xlsx"
        wb.save(p)
        with pytest.raises(ValueError, match="more than once"):
            read_labeled_matrix(str(p), "SAM")


def test_sut_readers_reject_repeated_codes(tmp_path):
    import openpyxl
    from test_data_modules import _make_sut_workbook, _make_use_workbook
    from edikit.data.sut import read_supply_table, read_use_table
    cases = [  # (builder, reader, edit)
        (_make_sut_workbook, read_supply_table,
         lambda ws: ws.cell(row=1, column=9, value="I1")),
        (_make_sut_workbook, read_supply_table,
         lambda ws: ws.append(["P2", "012", "Fruit", 62.0, 2.0, 10.0,
                               50.0, 50.0, 0.0])),
        (_make_use_workbook, read_use_table,
         lambda ws: ws.cell(row=1, column=7, value="I1")),
        (_make_use_workbook, read_use_table,
         lambda ws: ws.append(["P1", "011", "Cereals", 130.0, 10.0, 40.0,
                               30.0, 70.0, 40.0, 20.0])),
    ]
    for n, (build, read, edit) in enumerate(cases):
        p = tmp_path / f"t{n}.xlsx"
        build(str(p))
        wb = openpyxl.load_workbook(p)
        edit(wb.active)
        wb.save(p)
        with pytest.raises(ValueError, match="more than once"):
            read(str(p), 2019)


# --- every kind of recorded finding is exercised --------------------------

def test_check_supply_table_reports_each_finding_type(tmp_path):
    from test_data_modules import _make_sut_workbook
    from edikit.data.sut import check_supply_table, read_supply_table
    p = tmp_path / "sut.xlsx"
    _make_sut_workbook(str(p))
    t = read_supply_table(str(p), 2019)
    del t.total_supply_basic["P2"]
    t.embedded_output_totals["I2"] = 999.0
    got = {(f.check, f.subject) for f in check_supply_table(t)}
    assert got == {("row-identity", "P2"), ("column-total", "I2")}


def test_check_use_table_reports_each_finding_type(tmp_path):
    from test_data_modules import _make_use_workbook
    from edikit.data.sut import check_use_table, read_use_table
    p = tmp_path / "use.xlsx"
    _make_use_workbook(str(p))
    t = read_use_table(str(p), 2019)
    t.total_supply_purchasers["P1"] = 999.0           # row identity broken
    del t.total_supply_purchasers["P2"]               # row total missing
    t.embedded_intermediate_by_industry["I1"] = 999.0
    t.value_added[("D1", "I2")] = 99.0                # B1 unchanged
    got = {(f.check, f.subject) for f in check_use_table(t)}
    assert got == {("use-row-identity", "P1"), ("use-row-identity", "P2"),
                   ("use-column-total", "I1"), ("va-identity", "I2")}


def test_check_tre_reports_each_finding_type(tmp_path):
    from test_data_modules import _make_tre_workbook
    from edikit.data.ansd import check_tre, read_tre
    p = tmp_path / "tre.xlsx"
    _make_tre_workbook(str(p))
    tre = read_tre(str(p), 2014)
    tre.supply_basic["A00"] = 999.0
    tre.supply_purchaser["B00"] = 999.0
    tre.make[("A00", "A00")] = 150.0                  # branch output 150
    tre.operating_surplus["B00"] = 999.0
    got = {(f.check, f.subject) for f in check_tre(tre)}
    assert got == {("supply-origin", "A00"), ("supply-use-balance", "B00"),
                   ("output-identity", "A00"), ("va-decomposition", "B00")}


def test_concordance_validate_reports_unknown_sources():
    from edikit.data.concordance import Concordance
    c = Concordance(name="t", mapping={"I1": "a1", "I2": "a1"})
    problems = c.validate(["I1"])
    assert len(problems) == 1 and "unknown sources: I2" in problems[0]


def test_generate_records_identity_findings_cross_diffs_and_suppression(
        tmp_path):
    """A broken value-added cell and a supply cell below the published total
    are recorded, not refused: identity findings, a supply-use
    cross-difference, and the detail's coverage of the total."""
    from conftest import synthetic_economy_paths
    from edikit.pipeline.eurostat_sam import generate
    paths = synthetic_economy_paths(
        tmp_path, supply_edits={("A01", "CPA_A01"): 90},
        use_edits={("A01", "B1G"): 75})
    res = generate("XX", 2019, paths)
    assert any(f.startswith("output identity A01") for f in res.findings)
    assert any(f.startswith("VA identity A01") for f in res.findings)
    assert res.cross_diffs == ["A01"]
    assert abs(res.coverage["industry detail vs published supply total"]
               - 290 / 300) < 1e-9


def test_audit_rounding_bound_ignores_diagonal_cells():
    """A diagonal cell enters an account's receipts and payments alike, so
    its rounding cannot unbalance the account."""
    from edikit.pipeline.audit import audit
    res = audit({("a", "a"): 5.0, ("a", "b"): 3.0, ("b", "a"): 3.0})
    assert res.rounding_bound == {"a": 1.0, "b": 1.0}


def test_balance_alternative_targets_converge(tmp_path):
    """The receipts and payments targets (the build recipe's --targets) are
    scaled to the mean variant's total and balance like it."""
    from conftest import synthetic_economy_paths
    from edikit.pipeline.eurostat_sam import balance, generate
    res = generate("XX", 2019, synthetic_economy_paths(tmp_path))
    totals = {}
    for targets in ("mean", "receipts", "payments"):
        bal, _ = balance(res, targets=targets)
        assert bal.converged
        totals[targets] = sum(bal.matrix.values())
    assert abs(totals["receipts"] - totals["mean"]) < 1e-6
    assert abs(totals["payments"] - totals["mean"]) < 1e-6
